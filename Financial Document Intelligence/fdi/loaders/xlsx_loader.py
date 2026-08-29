import openpyxl

from fdi.schema import Document, Table


def load_xlsx(path: str, category: str) -> Document:
    """Extract each worksheet as a table into a normalized Document."""
    workbook = openpyxl.load_workbook(path, data_only=True)

    tables: list[Table] = []
    for sheet in workbook.worksheets:
        rows: Table = [
            [str(cell) if cell is not None else "" for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        tables.append(rows)

    return Document(
        source_path=path,
        category=category,
        file_type="xlsx",
        text="",
        tables=tables,
    )
